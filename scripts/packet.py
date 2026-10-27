"""작성용 자료 묶음 — 키워드 번호를 주면 키워드 정보 + 사례 원문 + 참고 원문 + 같은 주 관련 글을 출력한다.

  python -I scripts/packet.py 1 2 3          → 각 번호의 자료
  python -I scripts/packet.py --assign       → data/assign.json (번호별 사례 원문 배정) 생성

사례 원문은 제품·업종 단어가 제목/본문에 많이 나오는 글 중 덜 쓰인 것부터 배정(805편을 1,000편에 고르게).
"""
import json, os, re, sys
from collections import Counter
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XL = os.path.join(ROOT, '..', 'region_kw_prj', 'keywords', '한성쇼케이스_블로그키워드_1000.xlsx')
CORPUS = [json.loads(l) for l in open(os.path.join(ROOT, 'data', 'corpus.jsonl'), encoding='utf-8')]
ALL_CORPUS = CORPUS
# 사례 후보: 쇼케이스·냉장 이야기가 있는 글만 (드라이아이스·음식 상식 같은 게시판 글 제외)
CORPUS = [c for c in CORPUS if sum((c['title'] + ' '.join(c['paras'])).count(w) for w in ('쇼케이스', '냉장', '냉동', '진열')) >= 5]
BYLOG = {str(c['logNo']): c for c in ALL_CORPUS}
ROWS = [dict(zip(['no', 'week', 'kw', 'title', 'cat', 'prod', 'angle', 'sub', 'kind'], r[:9]))
        for r in load_workbook(XL, read_only=True)['키워드'].iter_rows(min_row=2, values_only=True)]
ASSIGN = os.path.join(ROOT, 'data', 'assign.json')
sys.stdout.reconfigure(encoding='utf-8')


def terms(r):
    t = {r['prod'], r['kw'].split()[0]} | {s.split()[0] for s in (r['sub'] or '').split(', ') if s}
    return {x for x in t if x and len(x) >= 2}


ANGLE_T = {'가격': '가격 견적 비용 금액', '견적': '견적 비용 가격 상담', '중고': '중고', '맞춤제작': '맞춤 제작 주문 출고', '추천': '추천 고르 선택',
           '온도': '온도 적정', '전기세': '전기 요금 전기세', '고장': '고장 수리 점검', 'AS': 'A/S AS 수리', '사이즈': '사이즈 규격 크기',
           '고르는법': '선택 고르 기준', '관리': '관리 청소', '업체': '업체 제작', '조명': '조명 LED', '결로': '결로 성에 김서림',
           '간냉식 직냉식': '간냉식 직냉식', '설치': '설치 출고 반입', '수명': '수명 오래', '렌탈': '렌탈 구매', '소음': '소음 소리',
           '디자인': '디자인 색상 인테리어 마감', '배치': '배치 동선 위치', '진열': '진열', '냉동겸용': '냉동', '도어': '도어 문',
           '콤프레셔': '콤프레셔 콤프', '여름철': '여름', '창업': '창업', '배송': '배송 반입 설치', '보증': 'A/S 보증 무상', '비교': '비교 차이'}


def score(c, ts, cat, angle=''):
    body = c['title'] + ' ' + ' '.join(c['paras'])
    at = ANGLE_T.get(angle, angle).split()
    return sum(min(body.count(x), 8) * (15 if x in c['title'] else 3) for x in ts) + (8 if c['cat'] == cat else 0) +         sum(min(body.count(x), 6) * (6 if x in c['title'] else 2) for x in at)


def assign():
    use = Counter(); out = {}
    keep = json.load(open(ASSIGN, encoding='utf-8')) if os.path.exists(ASSIGN) else {}
    import glob
    written = {int(os.path.basename(f)[:-5]) for f in glob.glob(os.path.join(ROOT, 'posts', '*.json'))}
    for r in ROWS:
        if r['no'] in written:  # 이미 쓴 글은 배정 유지
            out[r['no']] = keep[str(r['no'])]; use[keep[str(r['no'])]['case']] += 1; continue
        ts = terms(r)
        ranked = sorted(CORPUS, key=lambda c: (-score(c, ts, r['cat'], r['angle']) // 10 * 10 + use[c['logNo']] * 6, -len(c['imgs'])))
        cand = [c for c in ranked if len(c['imgs']) >= 5 and len(''.join(c['paras'])) > 600][:60]
        best = max(cand[:20], key=lambda c: score(c, ts, r['cat'], r['angle']) - use[c['logNo']] * 25)
        refs = [c['logNo'] for c in cand if c['logNo'] != best['logNo']][:2]
        use[best['logNo']] += 1
        out[r['no']] = dict(case=best['logNo'], refs=refs)
    json.dump(out, open(ASSIGN, 'w', encoding='utf-8'), ensure_ascii=False)
    print('assigned', len(out), 'distinct cases', len(set(v['case'] for v in out.values())), 'max reuse', max(use.values()))


def show(no):
    a = json.load(open(ASSIGN, encoding='utf-8'))[str(no)]
    r = next(x for x in ROWS if x['no'] == no)
    same = [x for x in ROWS if x['week'] == r['week'] and x['no'] != no and (x['cat'] == r['cat'] or x['prod'] == r['prod'])][:6]
    print(f"=== #{no} 주차 {r['week']} | 키워드: {r['kw']} | 각도: {r['angle']} | 제품·업종: {r['prod']} | 카테고리: {r['cat']}")
    print(f"제목안(참고만): {r['title']} | 보조 키워드: {r['sub']}")
    print('관련 글 후보(내부 링크용 번호): ' + ', '.join(f"#{x['no']} {x['kw']}" for x in same))
    for tag, lg in [('사례 원문', a['case'])] + [('참고 원문', x) for x in a['refs']]:
        c = BYLOG[str(lg)]
        body = '\n'.join(c['paras'])
        body = body[:600] if tag == '참고 원문' else body[:2200]
        print(f"\n--- {tag} logNo={lg} | {c['title']} | {c['cat']}")
        print(body)
        for tb in c['tables']:
            print('[표] ' + ' / '.join(' | '.join(row) for row in tb))


if __name__ == '__main__':
    if sys.argv[1] == '--assign':
        assign()
    else:
        for n in sys.argv[1:]:
            show(int(n)); print('\n' + '=' * 60 + '\n')
